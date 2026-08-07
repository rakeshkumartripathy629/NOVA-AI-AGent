"""
File processing service: text extraction, metadata enrichment and embedding.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional
from uuid import UUID

from app.core.logging import get_logger
from app.core.storage import storage_service
from app.models.file import FileStatus, FileType

logger = get_logger("services.files")


def _extract_text(filename: str, mime_type: str, content: bytes) -> str:
    """Extract plain text from file content based on type."""
    lower = filename.lower()

    if mime_type.startswith("text/") or lower.endswith((".md", ".markdown", ".csv", ".log", ".txt")):
        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return content.decode(encoding)
            except (UnicodeDecodeError, LookupError):
                continue
        return content.decode("utf-8", errors="replace")

    if lower.endswith(".pdf"):
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
            if text.strip():
                return text
        except Exception as exc:  # noqa: BLE001
            logger.warning("PDF extraction failed: %s", exc)

        try:
            import fitz 

            ocr_texts = []
            doc = fitz.open(stream=content, filetype="pdf")
            try:
                import easyocr

                reader = easyocr.Reader(["en"], verbose=False)
                for page in doc:
                    pix = page.get_pixmap(dpi=200)
                    img_bytes = pix.tobytes("png")
                    result = reader.readtext(img_bytes, detail=0, paragraph=True)
                    if result:
                        ocr_texts.append("\n".join(result))
            except Exception as exc:  # noqa: BLE001
                logger.warning("OCR failed: %s", exc)
            doc.close()
            if ocr_texts:
                return "\n\n".join(ocr_texts)
        except Exception as exc:  # noqa: BLE001
            logger.warning("PDF OCR fallback failed: %s", exc)

    if lower.endswith((".docx", ".doc")):
        try:
            from docx import Document

            doc = Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception as exc:  # noqa: BLE001
            logger.warning("DOCX extraction failed: %s", exc)

    if lower.endswith((".xlsx", ".xls")):
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines: List[str] = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(", ".join(cells))
            return "\n".join(lines)
        except Exception as exc:  # noqa: BLE001
            logger.warning("XLSX extraction failed: %s", exc)

    if lower.endswith((".json", ".jsonl")):
        try:
            text = content.decode("utf-8", errors="replace")
            import json

            return json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except Exception:  # noqa: BLE001
            return content.decode("utf-8", errors="replace")

    for ext in (".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css", ".sql", ".sh", ".java", ".go", ".rs", ".c", ".cpp", ".cs", ".rb", ".php", ".yaml", ".yml", ".toml", ".xml"):
        if lower.endswith(ext):
            return content.decode("utf-8", errors="replace")

    return ""


def _guess_file_type(filename: str, mime_type: str) -> FileType:
    lower = filename.lower()
    if mime_type.startswith("image/"):
        return FileType.IMAGE
    if mime_type.startswith("audio/"):
        return FileType.AUDIO
    if mime_type.startswith("video/"):
        return FileType.VIDEO
    if lower.endswith(".pdf"):
        return FileType.PDF
    if lower.endswith((".xlsx", ".xls", ".csv")):
        return FileType.SPREADSHEET
    if lower.endswith((".pptx", ".ppt")):
        return FileType.PRESENTATION
    if lower.endswith((".zip", ".tar", ".gz", ".rar", ".7z")):
        return FileType.ARCHIVE
    if lower.endswith((".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css", ".sql", ".sh", ".java", ".go", ".rs", ".c", ".cpp", ".cs", ".rb", ".php", ".yaml", ".yml", ".toml", ".xml")):
        return FileType.CODE
    if lower.endswith((".docx", ".doc", ".odt", ".md", ".markdown", ".txt", ".rtf")):
        return FileType.DOCUMENT
    if mime_type.startswith("text/"):
        return FileType.TEXT
    return FileType.OTHER


async def process_file(file_id: str, organization_id: str) -> Dict[str, Any]:
    """Download, extract and index an uploaded file."""
    from sqlalchemy import select

    from app.ai.rag import delete_document_chunks, index_document
    from app.db.session import get_session_factory
    from app.models.file import File

    session_factory = get_session_factory()
    async with session_factory() as db:
        result = await db.execute(select(File).where(File.id == UUID(file_id)))
        file_obj = result.scalar_one_or_none()
        if not file_obj:
            return {"status": "missing", "file_id": file_id}

        if file_obj.status == FileStatus.READY:
            return {"status": "completed", "file_id": file_id}

        file_obj.status = FileStatus.PROCESSING
        await db.commit()

        try:
            content = await storage_service.download_file(file_obj.storage_path)
        except Exception as exc:  # noqa: BLE001
            file_obj.status = FileStatus.FAILED
            file_obj.processing_error = f"Download failed: {exc}"
            await db.commit()
            return {"status": "failed", "file_id": file_id, "error": file_obj.processing_error}

        text = _extract_text(file_obj.filename, file_obj.mime_type, content)
        file_obj.file_type = _guess_file_type(file_obj.filename, file_obj.mime_type)
        file_obj.extracted_text = text[: 1_000_000]
        file_obj.pages = None

        chunk_count = 0
        if text.strip() and file_obj.knowledge_base_id:
            try:
                chunk_count = await index_document(
                    knowledge_base_id=file_obj.knowledge_base_id,
                    document_id=file_obj.id,
                    title=file_obj.original_filename,
                    content=text,
                    source_type="file",
                    organization_id=UUID(organization_id),
                )
            except Exception as exc:  # noqa: BLE001
                logger.warning("Failed to index file into knowledge base: %s", exc)
                try:
                    await delete_document_chunks(file_obj.id)
                except Exception:  # noqa: BLE001
                    pass

        file_obj.status = FileStatus.READY
        await db.commit()

        return {
            "status": "completed",
            "file_id": file_id,
            "file_type": file_obj.file_type.value if hasattr(file_obj.file_type, "value") else file_obj.file_type,
            "text_length": len(text),
            "chunk_count": chunk_count,
        }
